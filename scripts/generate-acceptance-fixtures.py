from __future__ import annotations

import argparse
import hashlib
import json
import mimetypes
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


MARKER = "AUTOMATED_SYNTHETIC_TEST_DATA"
SCHOOL = "星河实验学校"
PERSON = "林小航"


def font_path() -> Path | None:
    candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\arial.ttf"),
    ]
    return next((path for path in candidates if path.exists()), None)


def fixture_text() -> str:
    return (
        f"{MARKER}\n"
        "科瑞特 AI 自动化验收素材\n\n"
        f"虚构学校：{SCHOOL}\n"
        f"虚构学生：{PERSON}\n"
        "年级：五年级\n"
        "主题：校园节水观察\n\n"
        "所有姓名、学校和学习记录均为自动化合成数据，不含真实学生信息。\n"
    )


def rtf_escape(value: str) -> str:
    chunks: list[str] = []
    for character in value:
        codepoint = ord(character)
        if character in "\\{}":
            chunks.append("\\" + character)
        elif character == "\n":
            chunks.append("\\par\n")
        elif 32 <= codepoint <= 126:
            chunks.append(character)
        else:
            signed = codepoint if codepoint < 32768 else codepoint - 65536
            chunks.append(f"\\u{signed}?")
    return "".join(chunks)


def write_text_files(out: Path) -> None:
    body = fixture_text()
    (out / "synthetic-reading.txt").write_text(body, encoding="utf-8")
    rtf = "{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Arial;}}\\f0\\fs24 " + rtf_escape(body) + "}"
    (out / "synthetic-handbook.doc").write_bytes(rtf.encode("ascii"))
    (out / "synthetic-course.md").write_text(
        f"# 科瑞特 AI 自动化验收课程\n\n`{MARKER}`\n\n"
        f"- 学校：{SCHOOL}（虚构）\n"
        f"- 学生：{PERSON}（虚构）\n"
        "- 任务：写出三条校园节水建议。\n\n"
        "本文件由验收脚本自动生成，仅用于测试。\n",
        encoding="utf-8",
    )
    (out / "synthetic-subtitles.webvtt").write_text(
        "WEBVTT\n\n"
        "00:00:00.000 --> 00:00:01.500\n"
        f"{MARKER}\n\n"
        "00:00:01.500 --> 00:00:03.000\n"
        f"{SCHOOL}（虚构）验收视频\n",
        encoding="utf-8",
    )


def write_docx(out: Path) -> Path:
    path = out / "synthetic-handbook.docx"
    document = Document()
    title_style_properties = document.styles["Title"].element.get_or_add_pPr()
    title_border = title_style_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_style_properties.remove(title_border)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("科瑞特 AI 自动化验收手册")
    title_run.font.name = "Microsoft YaHei"
    title_run.font.size = Pt(24)
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    marker = document.add_paragraph()
    marker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    marker_run = marker.add_run(MARKER)
    marker_run.bold = True
    marker_run.font.color.rgb = RGBColor(160, 30, 30)
    document.add_heading("测试范围", level=1)
    document.add_paragraph(
        f"本手册用于验证课件上传、解析、预览和权限控制。"
        f"文中的{SCHOOL}和{PERSON}均为虚构数据。"
    )
    document.add_heading("课堂任务", level=1)
    document.add_paragraph("观察校园用水场景，记录一个问题，并提出可验证的改进方法。")
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "合成值"
    table.rows[0].cells[2].text = "用途"
    for values in [
        ("学校", SCHOOL, "权限与水印测试"),
        ("学生", PERSON, "解析与检索测试"),
        ("标记", MARKER, "防止误用为真实数据"),
    ]:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    document.save(path)
    return path


def write_pdf(out: Path) -> Path:
    path = out / "synthetic-worksheet.pdf"
    family = "Helvetica"
    selected_font = font_path()
    if selected_font:
        family = "FixtureFont"
        pdfmetrics.registerFont(TTFont(family, str(selected_font), subfontIndex=0))
    page = canvas.Canvas(str(path), pagesize=letter)
    page.setTitle("科瑞特 AI 自动化验收练习单")
    page.setFont(family, 20)
    page.drawString(72, 720, "科瑞特 AI 自动化验收练习单")
    page.setFont(family, 10)
    page.setFillColorRGB(0.65, 0.1, 0.1)
    page.drawString(72, 690, MARKER)
    page.setFillColorRGB(0, 0, 0)
    page.setFont(family, 12)
    lines = [
        f"虚构学校：{SCHOOL}",
        f"虚构学生：{PERSON}",
        "任务：列出两个校园节水观察点。",
        "声明：本文件仅用于自动化测试。",
    ]
    for index, line in enumerate(lines):
        page.drawString(72, 650 - index * 34, line)
    page.showPage()
    page.save()
    return path


def write_images(out: Path) -> Path:
    source = out / "synthetic-cover.png"
    image = Image.new("RGB", (1280, 720), "#EAF4FF")
    draw = ImageDraw.Draw(image)
    selected_font = font_path()
    title_font = ImageFont.truetype(str(selected_font), 58) if selected_font else ImageFont.load_default()
    body_font = ImageFont.truetype(str(selected_font), 30) if selected_font else ImageFont.load_default()
    draw.rectangle((80, 80, 1200, 640), outline="#1F5D8F", width=8)
    draw.text((120, 150), "科瑞特 AI 自动化验收封面", fill="#17324D", font=title_font)
    draw.text((120, 270), MARKER, fill="#A01E1E", font=body_font)
    draw.text((120, 350), f"{SCHOOL}（虚构）", fill="#17324D", font=body_font)
    draw.text((120, 410), f"{PERSON}（虚构）", fill="#17324D", font=body_font)
    image.save(source)
    image.save(out / "synthetic-cover.jpg", quality=92, optimize=True)
    image.save(out / "synthetic-cover.webp", quality=90, method=6)
    return source


def write_videos(out: Path, image_path: Path, ffmpeg: Path | None) -> list[Path]:
    if ffmpeg is None:
        return []
    outputs = [
        (out / "synthetic-lesson.mp4", ["-c:v", "libx264", "-pix_fmt", "yuv420p"]),
        (out / "synthetic-lesson.webm", ["-c:v", "libvpx-vp9", "-pix_fmt", "yuv420p"]),
    ]
    for output, codec in outputs:
        subprocess.run(
            [
                str(ffmpeg), "-y", "-loop", "1", "-i", str(image_path),
                "-t", "3", "-r", "24", *codec, "-an", str(output),
            ],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.PIPE,
        )
    return [item[0] for item in outputs]


def manifest(out: Path) -> None:
    rows = []
    for path in sorted(out.iterdir()):
        if not path.is_file() or path.name == "manifest.json" or path.name.startswith("~$"):
            continue
        payload = path.read_bytes()
        rows.append(
            {
                "file": path.name,
                "bytes": len(payload),
                "sha256": hashlib.sha256(payload).hexdigest(),
                "mime": {
                    ".webp": "image/webp",
                    ".webvtt": "text/vtt",
                }.get(path.suffix.lower(), mimetypes.guess_type(path.name)[0] or "application/octet-stream"),
                "classification": MARKER,
            }
        )
    (out / "manifest.json").write_text(
        json.dumps({"classification": MARKER, "files": rows}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate synthetic acceptance fixtures.")
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--ffmpeg", type=Path)
    parser.add_argument("--manifest-only", action="store_true")
    args = parser.parse_args()
    out = args.output_dir.resolve()
    out.mkdir(parents=True, exist_ok=True)
    if not args.manifest_only:
        write_text_files(out)
        write_docx(out)
        write_pdf(out)
        image_path = write_images(out)
        write_videos(out, image_path, args.ffmpeg)
    manifest(out)


if __name__ == "__main__":
    main()
